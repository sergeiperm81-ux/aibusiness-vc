from datetime import date
from pathlib import Path

from docx import Document


def build_manual_docx(output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    today = date.today().isoformat()

    doc.add_heading("Super.tennis - Manual Implementation Guide", level=0)
    doc.add_paragraph(f"Language: English only | Date: {today} | Package: Standard (EUR 149)")

    doc.add_heading("Why this document exists", level=1)
    doc.add_paragraph(
        "This guide is a fallback path for teams that prefer manual execution. "
        "It explains what to change, why each change matters, and how to verify results "
        "without relying on AI-assisted coding."
    )

    doc.add_heading("Who should use it", level=1)
    doc.add_paragraph(
        "Use this guide if implementation is handled by a developer, SEO manager, or content team "
        "working directly in the CMS or codebase."
    )

    doc.add_heading("What success looks like", level=1)
    doc.add_paragraph("Target outcome after implementation:")
    for item in [
        "Schema markup score improves from 65 to 85+.",
        "Content structure score improves from 78 to 88+.",
        "Citation readiness improves from 86 to 92+.",
        "Overall AI visibility score lands in the 93-97 range.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Pre-implementation checklist", level=1)
    for item in [
        "Create a backup branch or staging snapshot.",
        "Export current metadata and key page templates.",
        "List priority page clusters: /gear, /players, /records, /tournaments, /vs.",
        "Assign one owner for technical edits and one owner for content edits.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Manual implementation plan", level=1)

    doc.add_heading("Step 1 - Schema rollout by page type", level=2)
    doc.add_paragraph(
        "Add or upgrade JSON-LD manually by template. Use one clean JSON-LD block per object."
    )
    for item in [
        "Home: Organization + WebSite.",
        "Article pages: Article + BreadcrumbList.",
        "FAQ sections: FAQPage.",
        "Category pages: BreadcrumbList.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "Validation rule: every JSON-LD block must pass syntax validation and match visible page content."
    )

    doc.add_heading("Step 2 - llms.txt refresh", level=2)
    doc.add_paragraph(
        "Update /llms.txt with current high-value URLs and short factual descriptions. "
        "Keep entries clear, stable, and non-promotional."
    )

    doc.add_heading("Step 3 - FAQ and extraction blocks", level=2)
    doc.add_paragraph(
        "On high-intent pages, add 4-6 concise FAQs and one short summary block under the first major section."
    )
    for item in [
        "Each FAQ answer should stay under 90 words.",
        "Do not invent claims that are not present in source content.",
        "Use consistent section formatting across similar templates.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Step 4 - Comparison formatting", level=2)
    doc.add_paragraph(
        "For decision-focused pages, add a fixed comparison table with columns: "
        "Option, Best for, Strength, Limitation, Price band."
    )

    doc.add_heading("Step 5 - Internal link graph", level=2)
    doc.add_paragraph(
        "Strengthen contextual links between hub and spoke pages."
    )
    for item in [
        "Each spoke page links back to the relevant hub.",
        "Each hub page links out to at least five related spokes.",
        "Anchor text must describe the destination topic.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Step 6 - Metadata normalization", level=2)
    doc.add_paragraph(
        "Rewrite weak meta descriptions to a consistent 130-160 character format: "
        "search intent + concrete value + page context."
    )

    doc.add_heading("Manual QA checklist", level=1)
    for item in [
        "No JSON-LD syntax errors.",
        "No duplicate H1 elements.",
        "No broken internal links.",
        "No contradictory page claims.",
        "FAQ blocks are relevant and non-generic.",
        "Comparison tables add real decision value.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Definition of done", level=1)
    doc.add_paragraph(
        "Implementation is complete when all P1 tasks are merged, QA checklist is fully passed, "
        "and a re-scan confirms measurable improvement in schema, content structure, and citation readiness."
    )

    doc.add_heading("Delivery map", level=1)
    doc.add_paragraph("Use the package in this order:")
    for item in [
        "Executive PDF: for leadership decision-making.",
        "This Word guide: for manual implementation teams.",
        "Ops Pack markdown files: for AI-assisted execution.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(str(output_path))


if __name__ == "__main__":
    out = Path("deliverables/super-tennis/client/SuperTennis-Manual-Implementation-Guide.docx")
    build_manual_docx(out)
    print(f"Generated: {out}")

